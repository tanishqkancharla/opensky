"""Real saved-file specimens for grading acceptance, never GUI acceptance.

Completed specimens are built from the original inputs, not copied from gold.
Every specimen and reference is independently exported by production LibreOffice.
The campaign scorer and historical artifacts are not changed by this diagnostic.
"""
import argparse
import json
import sys
import zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[2] / "evals/parity/osworld"
sys.path.insert(0, str(ROOT))
from audit_impress_export import audit, digest, export_copy
from score_extended import score_extended


def duplicate_final_slides(source, destination):
    # Copy actual slide parts and their relationships, preserving media/styles.
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
    presentation = etree.fromstring(parts["ppt/presentation.xml"])
    relationships = etree.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    slides = presentation.find(f"{{{pns}}}sldIdLst")
    targets = {rel.get("Id"): rel.get("Target") for rel in relationships}
    final_two = list(slides)[-2:]
    next_id = max(int(slide.get("id")) for slide in slides) + 1
    for offset, slide in enumerate(final_two):
        old = targets[slide.get(f"{{{rns}}}id")]
        target = f"slides/opensky-control-{offset}.xml"
        rid = f"openskyControl{offset}"
        parts[f"ppt/{target}"] = parts[f"ppt/{old}"]
        old_rels = f"ppt/slides/_rels/{Path(old).name}.rels"
        if old_rels in parts:
            parts[f"ppt/slides/_rels/{Path(target).name}.rels"] = parts[old_rels]
        etree.SubElement(relationships, f"{{{relns}}}Relationship", Id=rid,
                         Type=f"{rns}/slide", Target=target)
        etree.SubElement(slides, f"{{{pns}}}sldId", id=str(next_id + offset),
                         attrib={f"{{{rns}}}id": rid})
        etree.SubElement(content_types, f"{{{ctns}}}Override", PartName=f"/ppt/{target}",
                         ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml")
    for name, tree in [("ppt/presentation.xml", presentation),
                       ("ppt/_rels/presentation.xml.rels", relationships),
                       ("[Content_Types].xml", content_types)]:
        parts[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def text_shapes(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from text_shapes(shape.shapes)
        elif shape.has_text_frame and shape.text:
            yield shape


def build_completed(task, source, output):
    key = task["id"].split("-")[0]
    if key == "9ec204e4":
        duplicate_final_slides(source, output)
        return
    deck = Presentation(source)
    if key == "9cf05d24":
        fill = deck.slides[0].background.fill
        fill.solid()
        # Preserve the upstream task's specific green; color tolerance is a
        # separate benchmark design issue, not silently changed here.
        fill.fore_color.rgb = RGBColor.from_string("00A933")
    elif key == "5cfb9197":
        table = next(shape.table for shape in deck.slides[3].shapes if shape.has_table)
        for column in range(4):
            runs = table.cell(0, column).text_frame.paragraphs[0].runs
            runs[0].text = f"T{column + 1}"
            for run in runs[1:]:
                run.text = ""
    elif key == "3161d64e":
        for shape, points in zip(text_shapes(deck.slides[13].shapes), [60, 28]):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(points)
    elif key == "2b94c692":
        image = next(shape for shape in deck.slides[1].shapes if shape.shape_type == 13)
        image.left = deck.slide_width - image.width
    else:
        raise ValueError(key)
    deck.save(output)


def build_incorrect(task, completed, output):
    deck = Presentation(completed)
    key = task["id"].split("-")[0]
    if key == "9cf05d24":
        deck.slides[0].background.fill.fore_color.rgb = RGBColor(255, 0, 0)
    elif key == "5cfb9197":
        table = next(shape.table for shape in deck.slides[3].shapes if shape.has_table)
        table.cell(0, 3).text_frame.paragraphs[0].runs[0].text = "T5"
    elif key == "9ec204e4":
        slides = deck.slides._sldIdLst
        slides.insert(len(slides) - 2, slides[-1])
    elif key == "3161d64e":
        for paragraph in list(text_shapes(deck.slides[13].shapes))[1].text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(22)
    elif key == "2b94c692":
        image = next(shape for shape in deck.slides[1].shapes if shape.shape_type == 13)
        image.left = 2_000_000  # Moved, but still short of the requested right side.
    deck.save(output)


def run(task_id, office, output):
    baseline = audit(task_id, office, output, reference_policy="task-preserving-export-v1")
    task = next(task for task in json.loads((ROOT / "manifest.json").read_text())["tasks"]
                if task["id"] == task_id)
    source = ROOT / task_id / task["inputFile"]
    specimens = output / "specimens"
    specimens.mkdir()
    completed = specimens / "completed.pptx"
    build_completed(task, source, completed)
    build_incorrect(task, completed, specimens / "incorrect.pptx")
    deck = Presentation(completed)
    # Unrelated text is an externally visible preservation requirement.
    shape = next(shape for slide in deck.slides for shape in text_shapes(slide.shapes))
    next(run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text).text += " UNREQUESTED CHANGE"
    deck.save(specimens / "damaged-content.pptx")
    deck = Presentation(completed)
    fill = deck.slides[-1].background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 0, 255)
    deck.save(specimens / "damaged-format.pptx")
    results = {"unchanged": baseline["unchangedTaskExportedReference"]}
    for kind in ["completed", "incorrect", "damaged-content", "damaged-format"]:
        exported = export_copy(office, specimens / f"{kind}.pptx", output / kind)
        before = digest(exported)
        results[kind] = score_extended(task, exported, reference_root=output / "exported-references")
        if digest(exported) != before:
            raise RuntimeError("Read-only grader changed a specimen")
    result = {"taskId": task_id, "diagnosticOnly": True, "modelSpend": 0,
              "referencePolicy": baseline["referencePolicy"],
              "office": baseline["office"], "controls": results,
              "cleanupVerified": all(json.loads(path.read_text())["temporaryProfileRemoved"] and
                                     json.loads(path.read_text())["ownedProcessGroupExited"]
                                     for path in output.rglob("*.export.json"))}
    (output / "controls.json").write_text(json.dumps(result, indent=2))
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--task", required=True)
    parser.add_argument("--libreoffice", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(run(args.task, args.libreoffice.resolve(), args.output.resolve()), indent=2))
