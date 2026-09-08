"""Explicit task-preserving reference adaptation for the duplicate-slide task.

The supplied gold deck changes unrelated earlier slides. Build the requested
A,B,A,B ending from the pinned input instead; retain the original gold score as
a separate result. This prepares a reference, never changes an agent artifact.
"""
from pathlib import Path
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlidePart

DUPLICATE_SLIDES_TASK = "9ec204e4-f0a3-42f8-8458-b772a6797cab"


def duplicate_slide_reference(source: Path, output: Path):
    deck = Presentation(source)
    originals = list(deck.slides)[-2:]
    for offset, slide in enumerate(originals):
        part = SlidePart.load(PackURI(f"/ppt/slides/opensky-reference-{offset}.xml"),
                              slide.part.content_type, deck.part.package, slide.part.blob)
        relationships = {}
        for rel in slide.part.rels.values():
            target = rel.target_ref if rel.is_external else rel.target_part
            relationships[rel.rId] = part.relate_to(target, rel.reltype, rel.is_external)
        namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for element in part._element.iter():
            for attribute, value in list(element.attrib.items()):
                if attribute.startswith(namespace) and value in relationships:
                    element.set(attribute, relationships[value])
        rid = deck.part.relate_to(part, RT.SLIDE)
        deck.slides._sldIdLst.add_sldId(rid)
    deck.save(output)
