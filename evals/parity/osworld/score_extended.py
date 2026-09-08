"""Original OSWorld file graders with separately reported content guards.

This module only reads saved files. No desktop actions, recalculation, repair,
or exporter is available to the grader.
"""
import contextlib
import io
import json
from pathlib import Path

from desktop_env.evaluators.metrics.table import compare_table
from desktop_env.evaluators.metrics.slides import compare_pptx_files
from desktop_env.evaluators.metrics.vscode_text import compare_text_file

ROOT = Path(__file__).resolve().parent
FUNCTIONS = {"compare_table": compare_table, "compare_pptx_files": compare_pptx_files,
             "compare_text_file": compare_text_file}


def workbook_data(path):
    from openpyxl import load_workbook
    book = load_workbook(path, data_only=True)
    # Sparse storage avoids enumerating a million blank validation cells.
    # Coordinates and saved cell values are the externally visible workbook.
    return [(sheet.title, {cell.coordinate: cell.value for cell in sheet._cells.values()
                          if cell.value is not None}) for sheet in book]


def slide_text(path):
    from pptx import Presentation
    def shapes_text(shapes):
        values = []
        for shape in shapes:
            if hasattr(shape, "shapes"):
                values.append(("group", shapes_text(shape.shapes)))
            elif shape.has_table:
                values.append(("table", [[cell.text for cell in row.cells] for row in shape.table.rows]))
            elif shape.has_text_frame:
                values.append(("text", shape.text))
        return values
    return [shapes_text(slide.shapes) for slide in Presentation(path).slides]


def content_matches(task, actual, expected):
    if task["category"] == "libreoffice_calc":
        return workbook_data(actual) == workbook_data(expected)
    if task["category"] == "libreoffice_impress":
        return slide_text(actual) == slide_text(expected)
    # An exact textual result is the user-visible outcome for the editor tasks.
    # Accept platform line endings, but preserve whitespace within each line.
    return Path(actual).read_text() == Path(expected).read_text()


def format_matches(task, actual, expected):
    if task["id"] == "12382c62-0cd1-4bf2-bdc8-1d20bf9b2371":
        from openpyxl import load_workbook
        def chart_kinds(path):
            book = load_workbook(path)
            return [(chart.tagname, chart.grouping, chart.barDir) for chart in book["Sheet2"]._charts]
        # The upstream type check does not distinguish clustered/stacked bars.
        return bool(chart_kinds(actual)) and chart_kinds(actual) == chart_kinds(expected)
    return True


def score_extended(task, output):
    upstream = json.loads((ROOT / task["id"] / "upstream-task.json").read_text())["evaluator"]
    functions = upstream["func"] if isinstance(upstream["func"], list) else [upstream["func"]]
    references = upstream["expected"] if isinstance(upstream["expected"], list) else [upstream["expected"]]
    options = upstream.get("options", {})
    options = options if isinstance(options, list) else [options] * len(functions)
    if len(functions) != len(references) or len(options) != len(functions) or any(f not in FUNCTIONS for f in functions):
        raise ValueError("Unsupported evaluator configuration; do not dispatch or score")
    if upstream.get("conj", "and") not in ["and", "or"]:
        raise ValueError("Unsupported evaluator conjunction")
    expected_paths = [ROOT / task["id"] / ref["dest"] for ref in references]
    # Reference/configuration faults are infrastructure errors, not task failures.
    for expected in expected_paths:
        if not expected.is_file():
            raise FileNotFoundError(expected)
    try:
        scores, content, formats = [], [], []
        for fn, expected, option in zip(functions, expected_paths, options):
            kwargs = dict(option)
            if fn == "compare_pptx_files":
                kwargs["enable_debug"] = False
            with contextlib.redirect_stdout(io.StringIO()):
                scores.append(FUNCTIONS[fn](str(output), str(expected), **kwargs))
            content.append(content_matches(task, output, expected))
            formats.append(format_matches(task, output, expected))
        aggregate = any if upstream.get("conj") == "or" else all
        passed = [s == 1 for s in scores]
        return {"upstreamScores": scores, "upstreamScore": int(aggregate(passed)),
                "contentMatches": aggregate(content), "formatMatches": aggregate(formats),
                "taskSuccess": aggregate([a and b and c for a, b, c in zip(passed, content, formats)])}
    except Exception as error:
        return {"upstreamScore": 0, "contentMatches": False, "formatMatches": False,
                "taskSuccess": False, "error": str(error)}
