"""Read only the agent's saved document; never repair or save it for the agent."""
import json
import sys
from pathlib import Path
from docx import Document
from document_metrics import compare_docx_files, is_first_line_centered

ROOT = Path(__file__).resolve().parent

def text(document):
    return {
        "paragraphs": [p.text for p in document.paragraphs],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
    }

def score(task_id, output):
    task = next(t for t in json.loads((ROOT / "manifest.json").read_text())["tasks"] if t["id"] == task_id)
    source = ROOT / task_id / task["inputFile"]
    try:
        actual = Document(output)
        if task["evaluator"] == "is_first_line_centered":
            upstream = is_first_line_centered(str(output))
            preserved = text(actual) == text(Document(source))
        else:
            gold = source.with_name(source.stem + "_Gold.docx")
            upstream = compare_docx_files(str(output), str(gold))
            # Upstream ignores table text; the user's all-text request includes it.
            preserved = text(actual) == text(Document(gold))
        return {"upstreamScore": upstream, "completeDocumentTextMatches": preserved, "taskSuccess": upstream == 1 and preserved}
    except Exception as error:
        return {"upstreamScore": 0, "completeDocumentTextMatches": False, "taskSuccess": False, "error": str(error)}

if __name__ == "__main__":
    print(json.dumps(score(sys.argv[1], Path(sys.argv[2]))))
